# -*- coding: utf-8 -*-
"""
Generador Maestro de la Presentación Ejecutiva de Clínica Cabal Salud SpA
Incluye:
- 15 diapositivas panorámicas 16:9 con diseño institucional y paleta corporativa
- Mapas conceptuales reales con conectores y verbos de enlace
- Comentarios y notas del orador exhaustivos en cada diapositiva para Pamela Quijada y Lilians Cortés
- Incorporación de todas las mejoras recientes (Asignar Sesión, consultante existente, domingo marcable, etc.)
"""
import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

def generar_presentacion_completa():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Paleta Oficial Cabal Salud
    C_NOCHE = RGBColor(16, 38, 60)       # #10263C (Azul Noche Oficial)
    C_PETROLEO = RGBColor(23, 75, 77)    # #174B4D (Verde Petróleo Oficial)
    C_GRANATE = RGBColor(119, 47, 59)    # #772F3B (Acento Granate Oficial)
    C_PERLA = RGBColor(247, 244, 237)    # #F7F4ED (Fondo Perla Oficial)
    C_BLANCO = RGBColor(255, 255, 255)
    C_ESMERALDA = RGBColor(16, 185, 129) # #10B981 (Aprobado/Activo)
    C_SLATE = RGBColor(100, 116, 139)    # #64748B
    C_PERLA_CARD = RGBColor(242, 238, 228)
    C_LIGHT_EMERALD = RGBColor(236, 253, 245)
    C_AMBER = RGBColor(217, 119, 6)

    def set_slide_background(slide, color=C_PERLA):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="CLÍNICA CABAL SALUD · ARQUITECTURA DIGITAL Y BLINDAJE LEGAL"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Calibri"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = C_PETROLEO
        p_cat.space_after = Pt(2)
        
        p_tit = tf.add_paragraph()
        p_tit.text = title_text
        p_tit.font.name = "Georgia"
        p_tit.font.size = Pt(21)
        p_tit.font.bold = True
        p_tit.font.color.rgb = C_NOCHE

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.48), Inches(11.7), Inches(0.04))
        line.fill.solid()
        line.fill.fore_color.rgb = C_PETROLEO
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=C_BLANCO, border_color=C_SLATE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1)
        else:
            card.line.fill.background()
        return card

    def add_concept_node(slide, left, top, width, height, title, subtitle=None, bg_color=C_BLANCO, border_color=C_PETROLEO, text_color=C_NOCHE):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.08)
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Georgia"
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = text_color
        p.alignment = PP_ALIGN.CENTER
        
        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.name = "Calibri"
            p2.font.size = Pt(9)
            p2.font.color.rgb = text_color
            p2.alignment = PP_ALIGN.CENTER
        return shape

    def add_link_pill(slide, left, top, width, height, text, bg_color=C_PERLA, text_color=C_PETROLEO):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = text_color
        shape.line.width = Pt(1)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.08)
        tf.margin_top = Inches(0.04)
        
        p = tf.paragraphs[0]
        p.text = text
        p.font.name = "Calibri"
        p.font.size = Pt(8.5)
        p.font.bold = True
        p.font.italic = True
        p.font.color.rgb = text_color
        p.alignment = PP_ALIGN.CENTER
        return shape

    def add_connector_line(slide, x1, y1, x2, y2, color=C_PETROLEO, width=1.5):
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
        conn.line.color.rgb = color
        conn.line.width = Pt(width)
        return conn

    def add_slide_notes(slide, text):
        notes_slide = slide.notes_slide
        tf = notes_slide.notes_text_frame
        tf.text = text

    # ==============================================================================
    # SLIDE 1: PORTADA EJECUTIVA
    # ==============================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1, C_NOCHE)

    decor = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(7.1), prs.slide_width, Inches(0.4))
    decor.fill.solid()
    decor.fill.fore_color.rgb = C_PETROLEO
    decor.line.fill.background()

    badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.1), Inches(4.2), Inches(0.45))
    badge.fill.solid()
    badge.fill.fore_color.rgb = C_PETROLEO
    badge.line.fill.background()
    b_tf = badge.text_frame
    b_tf.margin_top = Inches(0.08)
    b_p = b_tf.paragraphs[0]
    b_p.text = "RED DE SALUD ONLINE · CHILE"
    b_p.font.name = "Calibri"
    b_p.font.size = Pt(11)
    b_p.font.bold = True
    b_p.font.color.rgb = C_BLANCO
    b_p.alignment = PP_ALIGN.CENTER

    t_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.3))
    ttf = t_box.text_frame
    ttf.word_wrap = True

    p1 = ttf.paragraphs[0]
    p1.text = "Plataforma Web, Portal Profesional, Fichas Clínicas y Conexión con Gentia"
    p1.font.name = "Georgia"
    p1.font.size = Pt(34)
    p1.font.bold = True
    p1.font.color.rgb = C_PERLA
    p1.space_after = Pt(12)

    p2 = ttf.add_paragraph()
    p2.text = "Mapas Conceptuales de Arquitectura, Blindaje Legal ante Auditorías (Ley 20.584) y Pagos Directos"
    p2.font.name = "Calibri"
    p2.font.size = Pt(17)
    p2.font.color.rgb = RGBColor(180, 205, 220)
    p2.space_after = Pt(24)

    m_card = add_card(s1, Inches(1.0), Inches(5.0), Inches(11.3), Inches(1.5), RGBColor(22, 48, 75), C_PETROLEO)
    m_tf = m_card.text_frame
    m_tf.word_wrap = True
    m_tf.margin_left = m_tf.margin_right = Inches(0.4)
    m_tf.margin_top = Inches(0.2)

    mp1 = m_tf.paragraphs[0]
    mp1.text = "Clínica Cabal Salud SpA (cabalsalud.cl) · Motor Clínico Gentia (gentia.cl)"
    mp1.font.name = "Georgia"
    mp1.font.size = Pt(15)
    mp1.font.bold = True
    mp1.font.color.rgb = C_BLANCO
    mp1.space_after = Pt(4)

    mp2 = m_tf.add_paragraph()
    mp2.text = "Dirección Clínica y Administración: Pamela Quijada y Lilians Cortés | Blindaje Tributario (SII) y Sanitario (Ley 20.584)"
    mp2.font.name = "Calibri"
    mp2.font.size = Pt(12)
    mp2.font.color.rgb = RGBColor(160, 220, 210)

    add_slide_notes(s1, (
        "Estimadas Pamela y Lilians:\n\n"
        "Esta presentación ejecutiva fue preparada especialmente para ustedes como una guía completa y transparente sobre cómo está concebida y construida la plataforma de Clínica Cabal Salud SpA.\n\n"
        "Aquí podrán revisar:\n"
        "1. El sitio web de pacientes (cabalsalud.cl) con el catálogo de 9 profesionales y reserva directa.\n"
        "2. El Portal de Profesionales con horarios (incluyendo día domingo), duración flexible de sesión (30 a 90 min) y fichas ad-hoc.\n"
        "3. El Módulo de Administración Central, diseñado para que ustedes puedan registrar consultantes, filtrarlos en cascada en 2 pasos (especialidad y profesional), asignar citas con videollamada WebRTC y despachar correos corporativos duales automáticos.\n"
        "4. El Expediente Multidisciplinario con el caso transversal de prueba (Constanza Morales), derivación interna, descarga de informes y empaquetado integral en un archivo .ZIP con certificación legal SHA-256.\n"
        "5. El blindaje tributario ante el SII (Art. 12 E N° 8) y sanitario ante el MINSAL (Ley 20.584, D.S. N° 41 y Ley 21.719).\n\n"
        "En cada lámina les hemos dejado notas detalladas como esta explicando qué representa cada pantalla y qué decisiones o feedback esperamos de su parte. No hay ningún apuro: tómense el tiempo que necesiten para mirar todo con calma. Cualquier observación de forma (textos, colores, fotos) o de fondo (procedimientos, especialidades) será bienvenida."
    ))

    # ==============================================================================
    # SLIDE 2: FUNDAMENTO TRIBUTARIO (SII) Y EXENCIÓN DE IVA
    # ==============================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2, C_PERLA)
    add_header(s2, "Fundamento Tributario: Regla Suprema de Aislamiento Financiero (SII)")

    c_izq = add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf_izq = c_izq.text_frame
    tf_izq.word_wrap = True
    tf_izq.margin_left = tf_izq.margin_right = Inches(0.35)
    tf_izq.margin_top = Inches(0.3)

    p = tf_izq.paragraphs[0]
    p.text = "🏛️ Ley de IVA y Exención Médica (Art. 12 E N° 8)"
    p.font.name = "Georgia"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_GRANATE
    p.space_after = Pt(10)

    p = tf_izq.add_paragraph()
    p.text = (
        "• Marco Tributario Chileno:\n"
        "Las consultas y atenciones médicas, psicológicas y de salud ambulatoria están legalmente exentas de IVA ante el SII.\n\n"
        "• Riesgo Evitado con Gentia:\n"
        "Si gentia.cl recaudara los fondos de los pacientes para luego transferirlos a Cabal Salud, "
        "Gentia actuaría como mandatario/comisionista mercantil. Esto obligaría a gravar la comisión con IVA del 19% o arriesgar fiscalizaciones de triangulación.\n\n"
        "• Solución Definitiva: Direct Merchant Routing\n"
        "El 100% del dinero entra de forma directa y exclusiva a las cuentas bancarias de Clínica Cabal Salud SpA "
        "mediante Transbank Webpay Plus y Mercado Pago propios. Gentia nunca toca un solo peso."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.color.rgb = C_NOCHE

    c_der = add_card(s2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf_der = c_der.text_frame
    tf_der.word_wrap = True
    tf_der.margin_left = tf_der.margin_right = Inches(0.35)
    tf_der.margin_top = Inches(0.3)

    p = tf_der.paragraphs[0]
    p.text = "📦 Dos Repositorios 100% Desacoplados"
    p.font.name = "Georgia"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_PETROLEO
    p.space_after = Pt(10)

    p = tf_der.add_paragraph()
    p.text = (
        "1. Repositorio cabalsalud.cl (Vercel Independiente):\n"
        "   - Frontend con diseño, logo y paleta oficial.\n"
        "   - Catálogo multidisciplinario con 9 especialistas.\n"
        "   - Portal de Profesionales privado en el mismo dominio.\n"
        "   - Cobro directo /api/pagos/webpay y /mercadopago.\n\n"
        "2. Repositorio gentia.cl (Motor Clínico Headless):\n"
        "   - Motor de agenda y bloqueo de horarios en vivo.\n"
        "   - Bóveda de Fichas Clínicas Cifradas (AES-256).\n"
        "   - Salas WebRTC de teleconsulta cifrada punto a punto.\n"
        "   - Panel de Administración y Liquidaciones para Pamela y Lilians."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(12)
    p.font.color.rgb = C_NOCHE

    add_slide_notes(s2, (
        "Propósito de esta lámina: Explicar el modelo financiero y el blindaje tributario ante el Servicio de Impuestos Internos (SII).\n\n"
        "Puntos clave para Pamela y Lilians:\n"
        "• En Chile, las prestaciones ambulatorias de salud están exentas de IVA bajo el Artículo 12 Letra E N° 8 de la Ley de IVA.\n"
        "• Para evitar cualquier riesgo fiscal de triangulación de fondos, la plataforma opera bajo la 'Regla Suprema de Aislamiento Financiero': el 100% del dinero pagado por los consultantes entra de manera directa y exclusiva a las cuentas bancarias de Clínica Cabal Salud SpA a través de sus propias pasarelas de pago (Transbank Webpay Plus y Mercado Pago Chile).\n"
        "• Gentia jamás intermedia, retiene ni toca el dinero. Gentia actúa como proveedor de tecnología y software médico (SaaS), lo que blinda a Cabal Salud de cualquier intento del SII de calificar la operación como comisionista o aplicar el 19% de IVA sobre las consultas.\n\n"
        "Feedback esperado: Confirmar que las cuentas bancarias configuradas para la recaudación corresponden a las cuentas corrientes de Clínica Cabal Salud SpA."
    ))

    # ==============================================================================
    # SLIDE 3: MAPA CONCEPTUAL 1 (MACRO) — ARQUITECTURA DEL ECOSISTEMA
    # ==============================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3, C_PERLA)
    add_header(s3, "Mapa Conceptual 1: Ecosistema Integral, Flujo de Fondos y Custodia Clínica")

    # LÍNEAS CONECTORAS
    add_connector_line(s3, Inches(4.7), Inches(2.1), Inches(4.3), Inches(2.1), C_GRANATE, 2.0)
    add_connector_line(s3, Inches(2.6), Inches(2.26), Inches(2.6), Inches(2.6), C_GRANATE, 2.0)
    add_connector_line(s3, Inches(2.6), Inches(3.7), Inches(2.6), Inches(4.0), C_GRANATE, 2.0)
    add_connector_line(s3, Inches(2.6), Inches(4.36), Inches(2.6), Inches(4.7), C_GRANATE, 2.0)

    add_connector_line(s3, Inches(8.6), Inches(2.1), Inches(8.9), Inches(2.1), C_PETROLEO, 2.0)
    add_connector_line(s3, Inches(10.65), Inches(2.26), Inches(10.65), Inches(2.6), C_PETROLEO, 2.0)
    add_connector_line(s3, Inches(10.65), Inches(3.7), Inches(10.65), Inches(4.0), C_PETROLEO, 2.0)
    add_connector_line(s3, Inches(10.65), Inches(4.36), Inches(10.65), Inches(4.7), C_ESMERALDA, 2.0)

    add_connector_line(s3, Inches(6.65), Inches(2.8), Inches(6.65), Inches(3.1), C_PETROLEO, 2.0)
    add_connector_line(s3, Inches(6.65), Inches(3.46), Inches(6.65), Inches(3.7), C_PETROLEO, 2.0)
    add_connector_line(s3, Inches(6.65), Inches(4.65), Inches(6.65), Inches(4.9), C_PETROLEO, 2.0)
    add_connector_line(s3, Inches(6.65), Inches(5.26), Inches(6.65), Inches(5.5), C_NOCHE, 2.0)

    # Nodos
    add_concept_node(s3, Inches(4.7), Inches(1.8), Inches(3.9), Inches(1.0),
                     "CLÍNICA CABAL SALUD SpA",
                     "Entidad Titular y Prestadora de Salud (cabalsalud.cl)",
                     C_NOCHE, C_PETROLEO, C_PERLA)

    add_link_pill(s3, Inches(1.3), Inches(1.9), Inches(3.0), Inches(0.36),
                  "recauda 100% directo sin intermediar", C_PERLA, C_GRANATE)
    add_concept_node(s3, Inches(0.8), Inches(2.6), Inches(3.6), Inches(1.1),
                     "Cuentas Bancarias Cabal Salud",
                     "Transbank Webpay Plus & Mercado Pago Chile\n(Fondos van directo a cuenta de la clínica)",
                     C_BLANCO, C_GRANATE, C_NOCHE)

    add_link_pill(s3, Inches(1.0), Inches(4.0), Inches(3.2), Inches(0.36),
                  "blinda exención tributaria ante", C_PERLA, C_GRANATE)
    add_concept_node(s3, Inches(0.8), Inches(4.7), Inches(3.6), Inches(1.0),
                     "Servicio de Impuestos Internos (SII)",
                     "Exención de IVA garantizada (Art. 12 Letra E N° 8)\nEmisión de Boletas Exentas sin comisiones",
                     C_PERLA_CARD, C_GRANATE, C_NOCHE)

    add_link_pill(s3, Inches(8.9), Inches(1.9), Inches(3.3), Inches(0.36),
                  "delega custodia clínica sensible a", C_PERLA, C_PETROLEO)
    add_concept_node(s3, Inches(8.8), Inches(2.6), Inches(3.7), Inches(1.1),
                     "Motor Clínico Headless Gentia",
                     "Bóveda Cifrada de Salud (gentia.cl)\nCifrado en reposo AES-256 e inmutabilidad",
                     C_BLANCO, C_PETROLEO, C_NOCHE)

    add_link_pill(s3, Inches(9.0), Inches(4.0), Inches(3.3), Inches(0.36),
                  "blinda ante fiscalización legal de", C_PERLA, C_PETROLEO)
    add_concept_node(s3, Inches(8.8), Inches(4.7), Inches(3.7), Inches(1.0),
                     "Superintendencia de Salud / Agencia Datos",
                     "Cumplimiento Ley N° 20.584, D.S. N° 41 (15 años)\ny Ley N° 21.719 (Protección Datos Sensibles)",
                     C_LIGHT_EMERALD, C_ESMERALDA, C_NOCHE)

    add_link_pill(s3, Inches(5.0), Inches(3.1), Inches(3.3), Inches(0.36),
                  "ofrece atención multidisciplinaria a", C_PERLA, C_PETROLEO)
    add_concept_node(s3, Inches(4.5), Inches(3.7), Inches(4.3), Inches(0.95),
                     "Pacientes y Consultantes",
                     "Catálogo 9 profesionales, validación RUT y reserva in-page",
                     C_BLANCO, C_PETROLEO, C_NOCHE)

    add_link_pill(s3, Inches(5.0), Inches(4.9), Inches(3.3), Inches(0.36),
                  "provee entorno privado de gestión a", C_PERLA, C_PETROLEO)
    add_concept_node(s3, Inches(4.2), Inches(5.5), Inches(4.9), Inches(1.1),
                     "Portal Profesional (cabalsalud.cl)",
                     "Edición de perfil, arancel, tiempo de sesión (30 a 90 min),\nFichas Clínicas en Bóveda Gentia y datos de transferencia",
                     C_BLANCO, C_NOCHE, C_NOCHE)

    add_slide_notes(s3, (
        "Propósito de esta lámina: Mostrar visualmente el mapa conceptual general de la arquitectura que une lo financiero, lo clínico y lo normativo.\n\n"
        "Puntos clave para Pamela y Lilians:\n"
        "• Al centro: Clínica Cabal Salud SpA como entidad titular prestadora de salud.\n"
        "• A la izquierda (Rojo/Granate): El circuito de recaudación bancaria directa que garantiza la exención de IVA ante el SII.\n"
        "• A la derecha (Verde Petróleo y Esmeralda): La delegación de la custodia de historias clínicas a la Bóveda Cifrada Gentia, asegurando que Cabal Salud cumpla la exigencia del D.S. N° 41 (conservar fichas por 15 años de forma inalterable) y la Ley N° 20.584 de Derechos del Paciente.\n"
        "• Abajo (Azul Noche): La interacción con pacientes que agendan en la web pública y el equipo de profesionales que opera en el Portal Profesional.\n\n"
        "Feedback esperado: ¿Les parece claro este mapa conceptual de relaciones? Pueden fijarse en si los términos y conceptos reflejan adecuadamente la dinámica de trabajo de la clínica."
    ))

    # ==============================================================================
    # SLIDE 4: SCREENSHOT SITIO WEB PÚBLICO (cabalsalud.cl)
    # ==============================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4, C_PERLA)
    add_header(s4, "Sitio Web Oficial de Cara al Paciente (cabalsalud.cl)")

    c_info4 = add_card(s4, Inches(0.8), Inches(1.8), Inches(4.8), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf4 = c_info4.text_frame
    tf4.word_wrap = True
    tf4.margin_left = tf4.margin_right = Inches(0.3)
    tf4.margin_top = Inches(0.25)

    p = tf4.paragraphs[0]
    p.text = "Experiencia Paciente: Claridad y Cercanía"
    p.font.name = "Georgia"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NOCHE
    p.space_after = Pt(8)

    p = tf4.add_paragraph()
    p.text = (
        "• Hero Section:\n"
        "  Título oficial en tipografía serif cálida:\n"
        "  'Profesionales de salud para acompañarte con cercanía y criterio.'\n\n"
        "• Catálogo Multidisciplinario:\n"
        "  9 profesionales en Psicología, Nutrición, Matronería, Terapia Ocupacional, "
        "Fonoaudiología, Trabajo Social, Psiquiatría y Mediación Familiar.\n\n"
        "• Subespecialidades Estrella Destacadas:\n"
        "  Distintivo en Granate (#772F3B) para Deporte, Neurodivergencia, Perinatal y Geriatría.\n\n"
        "• Agendamiento In-Page:\n"
        "  Modal fluido sin salir del sitio para reservar y pagar."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(11.5)

    img_path_1 = os.path.join(os.getcwd(), "presentation_assets", "01_home.png")
    if os.path.exists(img_path_1):
        frame = add_card(s4, Inches(5.85), Inches(1.8), Inches(6.65), Inches(5.1), C_BLANCO, C_PETROLEO)
        s4.shapes.add_picture(img_path_1, Inches(5.95), Inches(1.9), width=Inches(6.45), height=Inches(4.6))
        cap_box = s4.shapes.add_textbox(Inches(5.95), Inches(6.55), Inches(6.45), Inches(0.3))
        cap_tf = cap_box.text_frame
        cap_p = cap_tf.paragraphs[0]
        cap_p.text = "Vista en vivo: https://cabalsalud.vercel.app (Diseño Perla, Azul Noche y Verde Petróleo)"
        cap_p.font.size = Pt(10)
        cap_p.font.color.rgb = C_SLATE
        cap_p.alignment = PP_ALIGN.CENTER

    add_slide_notes(s4, (
        "Propósito de esta lámina: Mostrar la imagen y experiencia que reciben los consultantes al visitar la página web de la clínica.\n\n"
        "Puntos clave para Pamela y Lilians:\n"
        "• La identidad visual aplica estrictamente la paleta corporativa: Fondo Perla (#F7F4ED), textos en Azul Noche (#10263C), acentos en Verde Petróleo (#174B4D) y etiquetas en Granate (#772F3B).\n"
        "• Catálogo multidisciplinario con los 9 profesionales que conforman la red, permitiendo explorar perfiles, subespecialidades (deporte, perinatal, neurodivergencia, etc.) y aranceles.\n"
        "• Agendamiento 'in-page': El paciente puede reservar su cita directamente en la página sin ser derivado a plataformas externas genéricas, lo que transmite mayor profesionalismo y seguridad al pagar.\n\n"
        "Feedback de forma y fondo: Les pedimos revisar si la redacción de los encabezados, las áreas de atención y la presentación de los profesionales reflejan con precisión la filosofía y el tono de Cabal Salud."
    ))

    # ==============================================================================
    # SLIDE 5: FLUJO OPERACIONAL PASO A PASO
    # ==============================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5, C_PERLA)
    add_header(s5, "Flujo Operacional: Paso a Paso de Agendamiento, Cobro y Confirmación")

    pasos = [
        ("PASO 1", "Selección en Vivo", "El paciente elige especialista y fecha. Consulta /api/agenda/disponibilidad que conecta en vivo con Gentia.", C_PETROLEO),
        ("PASO 2", "Validación RUT y Modalidad", "Formulario valida RUT Módulo 11 (para boleta de reembolso). Si elige 'Visita a Terreno' se exige dirección física.", C_NOCHE),
        ("PASO 3", "Pre-Agendamiento TTL", "Cabal Salud llama a Gentia /clinica/preAgendarCita. El slot queda bloqueado por 15 min en 'pendiente_pago'.", C_PETROLEO),
        ("PASO 4", "Transacción Bancaria", "Se inicia el cobro en Transbank Webpay Plus o Mercado Pago con credenciales exclusivas de Cabal Salud SpA.", C_GRANATE),
        ("PASO 5", "Retorno y Webhook HMAC", "Tras el pago exitoso, Cabal Salud confirma retorno y dispara Webhook firmado con SHA256 hacia Gentia.", C_ESMERALDA),
        ("PASO 6", "Confirmación y Calendarios", "Se activa la cita, se genera enlace WebRTC, se envía correo y se descarga .ICS / 1-clic Google Calendar.", C_PETROLEO)
    ]

    for idx, (step_num, step_title, step_desc, step_color) in enumerate(pasos):
        row = idx // 3
        col = idx % 3
        left_x = Inches(0.8 + col * 4.0)
        top_y = Inches(1.9 + row * 2.6)
        
        card = add_card(s5, left_x, top_y, Inches(3.7), Inches(2.35), C_BLANCO, step_color)
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.25)
        ctf.margin_top = Inches(0.2)
        
        p = ctf.paragraphs[0]
        p.text = f"{step_num} · {step_title}"
        p.font.name = "Georgia"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = step_color
        p.space_after = Pt(6)
        
        p = ctf.add_paragraph()
        p.text = step_desc
        p.font.name = "Calibri"
        p.font.size = Pt(11)
        p.font.color.rgb = C_NOCHE

    add_slide_notes(s5, (
        "Propósito de esta lámina: Detallar cronológicamente el circuito de 6 pasos que se ejecuta automáticamente cada vez que un paciente solicita una hora.\n\n"
        "Puntos clave para Pamela y Lilians:\n"
        "• Paso 1: El paciente visualiza en tiempo real los días y horas disponibles de cada especialista.\n"
        "• Paso 2: El sistema valida que el RUT ingresado sea auténtico mediante el algoritmo chileno de Módulo 11, evitando errores en la emisión de boletas para reembolso Isapre/FONASA.\n"
        "• Paso 3: Al avanzar al pago, el horario se bloquea durante 15 minutos (TTL) para que nadie más pueda ocuparlo simultáneamente.\n"
        "• Paso 4: El pago se procesa en Transbank o Mercado Pago, acreditándose de inmediato en la cuenta corriente de Cabal Salud.\n"
        "• Paso 5: El servidor de Cabal Salud notifica al motor de Gentia mediante una firma criptográfica privada (Webhook HMAC).\n"
        "• Paso 6: La cita queda confirmada, se crea el enlace WebRTC y el paciente recibe la confirmación con opción de agregar a Google Calendar o Outlook (.ICS).\n\n"
        "Feedback esperado: ¿Les parece adecuada la ventana de 15 minutos para que el paciente complete el pago antes de liberar el cupo?"
    ))

    # ==============================================================================
    # SLIDE 6: SCREENSHOT COMPROBANTE DE PAGO Y CONFIRMACIÓN
    # ==============================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6, C_PERLA)
    add_header(s6, "Comprobante Oficial y Sincronización de Calendarios")

    c_info6 = add_card(s6, Inches(0.8), Inches(1.8), Inches(4.8), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf6 = c_info6.text_frame
    tf6.word_wrap = True
    tf6.margin_left = tf6.margin_right = Inches(0.3)
    tf6.margin_top = Inches(0.25)

    p = tf6.paragraphs[0]
    p.text = "Respaldo y Resiliencia Stateless"
    p.font.name = "Georgia"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NOCHE
    p.space_after = Pt(8)

    p = tf6.add_paragraph()
    p.text = (
        "• Comprobante Inmutable de Pago:\n"
        "  Muestra Orden de Compra, Código de Autorización Transbank, 4 últimos dígitos de tarjeta y monto en CLP.\n\n"
        "• Datos Listos para Reembolso:\n"
        "  RUT del paciente, previsión (Isapre / FONASA) y datos del profesional con RNPI para presentar a reembolso médico.\n\n"
        "• Sala de Teleconsulta WebRTC:\n"
        "  Botón directo para ingresar a la sala cifrada de videollamada de Gentia sin requerir cuentas externas.\n\n"
        "• Integración 1-Clic de Calendario:\n"
        "  Botón nativo para Google Calendar y descarga de archivo estándar .ICS (Outlook, Apple Calendar)."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(11.5)

    img_path_4 = os.path.join(os.getcwd(), "presentation_assets", "04_confirmacion_reserva.png")
    if os.path.exists(img_path_4):
        frame = add_card(s6, Inches(5.85), Inches(1.8), Inches(6.65), Inches(5.1), C_BLANCO, C_PETROLEO)
        s6.shapes.add_picture(img_path_4, Inches(5.95), Inches(1.9), width=Inches(6.45), height=Inches(4.6))
        cap_box = s6.shapes.add_textbox(Inches(5.95), Inches(6.55), Inches(6.45), Inches(0.3))
        cap_tf = cap_box.text_frame
        cap_p = cap_tf.paragraphs[0]
        cap_p.text = "Vista en vivo: /confirmacion-reserva.html con voucher directo de Cabal Salud SpA"
        cap_p.font.size = Pt(10)
        cap_p.font.color.rgb = C_SLATE
        cap_p.alignment = PP_ALIGN.CENTER

    add_slide_notes(s6, (
        "Propósito de esta lámina: Mostrar la pantalla y comprobante de confirmación que recibe el paciente de forma inmediata tras completar su pago.\n\n"
        "Puntos clave para Pamela y Lilians:\n"
        "• Voucher con validez tributaria y médica: Incluye el código de autorización de Transbank, el monto cancelado en pesos chilenos y el RUT del paciente.\n"
        "• Datos para reembolso en Isapres: Muestra el nombre y número de Registro Nacional de Prestadores Individuales (RNPI) del especialista, requisito indispensable para que el paciente solicite la cobertura de su seguro.\n"
        "• Acceso directo a la sala WebRTC: Un botón claro para conectarse a la sesión online sin necesidad de descargar programas ni crear cuentas de terceros.\n"
        "• Sincronización con calendarios: Botones directos para agendar en Google Calendar o descargar el archivo .ICS (compatible con Outlook y iPhone), disminuyendo olvidos e inasistencias.\n\n"
        "Feedback esperado: ¿Desean añadir alguna recomendación particular en el comprobante, por ejemplo, indicaciones previas sobre el lugar o conexión para la sesión?"
    ))

    # ==============================================================================
    # SLIDE 7: MAPA CONCEPTUAL 2 — BLINDAJE LEGAL ANTE AUDITORÍAS Y FICHA CLÍNICA
    # ==============================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7, C_PERLA)
    add_header(s7, "Mapa Conceptual 2: Fichas Clínicas, Epicrisis y Blindaje ante Auditorías")

    # LÍNEAS CONECTORAS
    add_connector_line(s7, Inches(4.5), Inches(2.18), Inches(4.2), Inches(2.18), C_PETROLEO, 2.0)
    add_connector_line(s7, Inches(2.55), Inches(2.36), Inches(2.55), Inches(2.6), C_PETROLEO, 2.0)
    add_connector_line(s7, Inches(2.55), Inches(3.7), Inches(2.55), Inches(4.0), C_PETROLEO, 2.0)
    add_connector_line(s7, Inches(2.55), Inches(4.36), Inches(2.55), Inches(4.7), C_GRANATE, 2.0)

    add_connector_line(s7, Inches(8.8), Inches(2.18), Inches(9.1), Inches(2.18), C_ESMERALDA, 2.0)
    add_connector_line(s7, Inches(10.75), Inches(2.36), Inches(10.75), Inches(2.6), C_ESMERALDA, 2.0)
    add_connector_line(s7, Inches(10.75), Inches(3.7), Inches(10.75), Inches(4.0), C_ESMERALDA, 2.0)
    add_connector_line(s7, Inches(10.75), Inches(4.36), Inches(10.75), Inches(4.7), C_ESMERALDA, 2.0)

    add_connector_line(s7, Inches(6.65), Inches(2.9), Inches(6.65), Inches(3.2), C_PETROLEO, 2.0)
    add_connector_line(s7, Inches(6.65), Inches(3.56), Inches(6.65), Inches(3.8), C_PETROLEO, 2.0)
    add_connector_line(s7, Inches(6.65), Inches(4.8), Inches(6.65), Inches(5.0), C_NOCHE, 2.0)
    add_connector_line(s7, Inches(6.65), Inches(5.36), Inches(6.65), Inches(5.6), C_NOCHE, 2.0)

    # Nodos
    add_concept_node(s7, Inches(4.5), Inches(1.8), Inches(4.3), Inches(1.1),
                     "Especialista (Psicólogo / Terapeuta)",
                     "Opera desde el Portal Profesional en cabalsalud.cl\n(Acceso sin Google con Demo o Credenciales)",
                     C_NOCHE, C_PETROLEO, C_PERLA)

    add_link_pill(s7, Inches(1.0), Inches(2.0), Inches(3.2), Inches(0.36),
                  "configura de forma autónoma", C_PERLA, C_PETROLEO)
    add_concept_node(s7, Inches(0.8), Inches(2.6), Inches(3.5), Inches(1.1),
                     "Arancel y Tiempo de Sesión",
                     "Duración modificable (30, 45, 50, 60, 75 o 90 min)\ny valor de consulta en CLP sin intervención externa",
                     C_BLANCO, C_PETROLEO, C_NOCHE)

    add_link_pill(s7, Inches(1.0), Inches(4.0), Inches(3.2), Inches(0.36),
                  "redacta notas de atención y", C_PERLA, C_PETROLEO)
    add_concept_node(s7, Inches(0.8), Inches(4.7), Inches(3.5), Inches(1.2),
                     "Epicrisis Oficial / Alta Médica",
                     "Resumen psicoterapéutico, diagnóstico CIE-11/DSM-5\ne informe de avance clínico para Isapres",
                     C_BLANCO, C_GRANATE, C_NOCHE)

    add_link_pill(s7, Inches(9.1), Inches(2.0), Inches(3.3), Inches(0.36),
                  "sincroniza datos sensibles con", C_PERLA, C_ESMERALDA)
    add_concept_node(s7, Inches(9.0), Inches(2.6), Inches(3.5), Inches(1.1),
                     "Bóveda Cifrada Gentia (EHR Vault)",
                     "Encriptación AES-256 en reposo y tránsito\nFirma inmutable por sesión con hash de auditoría",
                     C_LIGHT_EMERALD, C_ESMERALDA, C_NOCHE)

    add_link_pill(s7, Inches(9.1), Inches(4.0), Inches(3.3), Inches(0.36),
                  "blinda legalmente a Cabal Salud ante", C_PERLA, C_ESMERALDA)
    add_concept_node(s7, Inches(9.0), Inches(4.7), Inches(3.5), Inches(1.2),
                     "Auditorías Sanitarias y de Datos",
                     "Ley N° 20.584 (Derechos del Paciente)\nD.S. N° 41 (Custodia 15 años) y Ley N° 21.719\nCabal Salud acredita custodia delegada",
                     C_BLANCO, C_ESMERALDA, C_NOCHE)

    add_link_pill(s7, Inches(5.0), Inches(3.2), Inches(3.3), Inches(0.36),
                  "atiende mediante sala segura a", C_PERLA, C_PETROLEO)
    add_concept_node(s7, Inches(4.5), Inches(3.8), Inches(4.3), Inches(1.0),
                     "Consultante con Consentimiento Firmado",
                     "Atención por videoconsulta WebRTC de Gentia\ncon consentimiento informado legalmente registrado",
                     C_BLANCO, C_PETROLEO, C_NOCHE)

    add_link_pill(s7, Inches(5.0), Inches(5.0), Inches(3.3), Inches(0.36),
                  "reporta honorarios atendidos a", C_PERLA, C_NOCHE)
    add_concept_node(s7, Inches(4.5), Inches(5.6), Inches(4.3), Inches(1.0),
                     "Liquidación Mensual (Pamela y Lilians)",
                     "Datos bancarios para transferencia directa de honorarios",
                     C_PERLA_CARD, C_NOCHE, C_NOCHE)

    add_slide_notes(s7, (
        "Propósito de esta lámina: Explicar el blindaje legal y la custodia documental de las historias clínicas ante exigencias sanitarias.\n\n"
        "Puntos clave para Pamela y Lilians:\n"
        "• Autonomía operativa del terapeuta: Cada profesional redacta sus evoluciones DAP y emite epicrisis clínicas para reembolso.\n"
        "• Custodia técnica en Bóveda Cifrada Gentia: Las historias clínicas no residen en el hosting web convencional, sino en una bóveda electrónica con cifrado AES-256 en reposo y tránsito.\n"
        "• Cumplimiento del D.S. N° 41 del MINSAL: El reglamento exige conservar las fichas clínicas por al menos 15 años bajo condiciones de integridad e inalterabilidad. Al contar con custodia delegada y sellado criptográfico inmutable, Cabal Salud acredita el cumplimiento estricto de la Ley N° 20.584 y la Ley N° 21.719 de Protección de Datos Personales.\n\n"
        "Feedback esperado: Validar este modelo de custodia delegada para presentar ante auditorías de la Superintendencia de Salud."
    ))

    # ==============================================================================
    # SLIDE 8: SCREENSHOT PORTAL DE PROFESIONALES (cabalsalud.cl)
    # ==============================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8, C_PERLA)
    add_header(s8, "Interfaz del Portal Profesional: Perfil, Aranceles y Tiempo de Sesión")

    c_info8 = add_card(s8, Inches(0.8), Inches(1.8), Inches(4.8), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf8 = c_info8.text_frame
    tf8.word_wrap = True
    tf8.margin_left = tf8.margin_right = Inches(0.3)
    tf8.margin_top = Inches(0.25)

    p = tf8.paragraphs[0]
    p.text = "Operación 100% en cabalsalud.cl"
    p.font.name = "Georgia"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_PETROLEO
    p.space_after = Pt(8)

    p = tf8.add_paragraph()
    p.text = (
        "• Sin depender de Google OAuth:\n"
        "  Selector de demostración para entrar con 1 clic como cualquiera de los 9 profesionales, o con correo/clave clínica.\n\n"
        "• Horarios y Día Domingo Configurable:\n"
        "  El domingo ahora es completamente marcable y editable en la grilla semanal de disponibilidad de cada terapeuta.\n\n"
        "• Duración de Sesión Modificable (30 a 90 min):\n"
        "  Cada especialista define si su sesión dura 30, 45, 50, 60, 75 o 90 minutos y su arancel en CLP.\n\n"
        "• Botones Claros y Selección Visible:\n"
        "  Botones simplificados ('Registrar' y 'Guardar') e indicador visual destacado al seleccionar un consultante en la lista.\n\n"
        "• Ficha Pública Previa y Transferencias:\n"
        "  Previsualizador de perfil público y registro de datos bancarios para transferencias directas de Pamela y Lilians."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(11)

    img_path_3 = os.path.join(os.getcwd(), "presentation_assets", "03_portal_profesional.png")
    if os.path.exists(img_path_3):
        frame = add_card(s8, Inches(5.85), Inches(1.8), Inches(6.65), Inches(5.1), C_BLANCO, C_PETROLEO)
        s8.shapes.add_picture(img_path_3, Inches(5.95), Inches(1.9), width=Inches(6.45), height=Inches(4.6))
        cap_box = s8.shapes.add_textbox(Inches(5.95), Inches(6.55), Inches(6.45), Inches(0.3))
        cap_tf = cap_box.text_frame
        cap_p = cap_tf.paragraphs[0]
        cap_p.text = "Vista en vivo: https://cabalsalud.vercel.app/portal-profesional.html"
        cap_p.font.size = Pt(10)
        cap_p.font.color.rgb = C_SLATE
        cap_p.alignment = PP_ALIGN.CENTER

    add_slide_notes(s8, (
        "Propósito de esta lámina: Presentar el espacio de trabajo que utilizarán los profesionales de Cabal Salud para gestionar sus consultas y disponibilidad.\n\n"
        "Mejoras recientes implementadas para Pamela y Lilians:\n"
        "• Día Domingo Habilitado: En la pestaña de 'Horarios y Disponibilidad', el domingo ahora es totalmente seleccionable y marcable para aquellos especialistas que deseen atender ese día.\n"
        "• Duración de sesión flexible: Cada terapeuta puede fijar autónomamente cuánto dura su consulta (30, 45, 50, 60, 75 o 90 minutos) y su valor en CLP.\n"
        "• Botones más sencillos: Se simplificaron los textos de los botones principales a 'Registrar' (para nuevas atenciones) y 'Guardar' (para epicrisis), dejando de lado términos técnicos complejos.\n"
        "• Selección visible de pacientes: Al hacer clic en un paciente del directorio, el botón cambia notablemente de color y resalta con un borde visible para saber siempre a quién se está atendiendo.\n"
        "• Datos bancarios para transferencias: Cada profesional registra su banco, cuenta corriente y RUT para que ustedes puedan realizar la liquidación de honorarios con rapidez.\n\n"
        "Feedback esperado: ¿Les parece cómodo el diseño del portal para que los terapeutas administren sus horarios y tarifas?"
    ))

    # ==============================================================================
    # SLIDE 9: NUEVA BÓVEDA DE FICHA CLÍNICA & EPICRISIS (LEY 20.584)
    # ==============================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9, C_PERLA)
    add_header(s9, "Gestión Clínica de Pacientes, Evoluciones DAP y Bóveda Gentia")

    c_info9 = add_card(s9, Inches(0.8), Inches(1.8), Inches(4.8), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf9 = c_info9.text_frame
    tf9.word_wrap = True
    tf9.margin_left = tf9.margin_right = Inches(0.3)
    tf9.margin_top = Inches(0.2)

    p = tf9.paragraphs[0]
    p.text = "Conectado al Motor de Gentia (Sin Migrar)"
    p.font.name = "Georgia"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_ESMERALDA
    p.space_after = Pt(6)

    p = tf9.add_paragraph()
    p.text = (
        "• Operación Clínica Completa en cabalsalud.cl:\n"
        "  Acceso ágil a 'Registrar Paciente', búsqueda en vivo por RUT/Nombre y sincronización total entre Próximas Sesiones y Fichas Clínicas.\n\n"
        "• Motivo Latente, CIE-11 y Banderas Rojas:\n"
        "  Registro de motivo manifiesto y motivo latente en psicología, antecedentes clínicos y alertas de riesgo vital (VIF, Suicidio).\n\n"
        "• Examen Mental V70 & Evoluciones DAP:\n"
        "  Evaluación de 8 esferas psíquicas, notas DAP con sliders de síntomas (Ansiedad, Depresión, Estrés 0-10) y sellado SHA-256.\n\n"
        "• Bóveda 15A Exclusiva de Administración:\n"
        "  La opción de archivo histórico 15A se reservó exclusivamente para la administración central de Cabal Salud.\n\n"
        "• Custodia Legal Delegada (15 Años):\n"
        "  Respaldo inmutable bajo AES-256 en Gentia (Ley N° 20.584, D.S. N° 41 y Ley N° 21.719)."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(10.5)

    img_path_6 = os.path.join(os.getcwd(), "presentation_assets", "06_ficha_clinica.png")
    if not os.path.exists(img_path_6):
        img_path_6 = os.path.join(os.getcwd(), "presentation_assets", "03_portal_profesional.png")

    if os.path.exists(img_path_6):
        frame = add_card(s9, Inches(5.85), Inches(1.8), Inches(6.65), Inches(5.1), C_BLANCO, C_ESMERALDA)
        s9.shapes.add_picture(img_path_6, Inches(5.95), Inches(1.9), width=Inches(6.45), height=Inches(4.6))
        cap_box = s9.shapes.add_textbox(Inches(5.95), Inches(6.55), Inches(6.45), Inches(0.3))
        cap_tf = cap_box.text_frame
        cap_p = cap_tf.paragraphs[0]
        cap_p.text = "Pestaña 'Ficha Clínica & Epicrisis' en el Portal Profesional de Cabal Salud"
        cap_p.font.size = Pt(10)
        cap_p.font.color.rgb = C_SLATE
        cap_p.alignment = PP_ALIGN.CENTER

    add_slide_notes(s9, (
        "Propósito de esta lámina: Mostrar cómo funciona la ficha clínica electrónica, las evoluciones de sesión y el resguardo de antecedentes.\n\n"
        "Ajustes aplicados según sus instrucciones:\n"
        "• Sincronización idéntica de pacientes: Tanto en la pestaña 'Próximas Sesiones' como en 'Fichas Clínicas', los datos de los pacientes asignados (incluyendo a la Ps. Camila Valdés y al resto del equipo) se reflejan de forma perfectamente coincidente.\n"
        "• Motivo de Consulta Manifiesto y Latente: En las fichas de psicología se incorporó formalmente el campo 'Motivo de Consulta Latente', permitiendo al terapeuta contrastar la demanda explícita con los factores subyacentes.\n"
        "• Bóveda 15A Exclusiva de Administración: Se retiró la opción de bóveda histórica 15A de la vista común de los terapeutas, manteniéndose exclusivamente como una herramienta directiva para Pamela y Lilians.\n"
        "• Evoluciones con notas DAP: Estructura profesional DAP (Datos, Análisis, Plan) con sliders visuales de 0 a 10 para evaluar la evolución cuantitativa de síntomas (ansiedad, estrés, ánimo) a lo largo del tratamiento.\n"
        "• Alertas críticas: Detección visual y registro especial de banderas rojas ante ideación suicida o violencia intrafamiliar.\n\n"
        "Feedback esperado: ¿Les parece acertada la inclusión del motivo latente y el formato de las notas DAP para la práctica de salud mental?"
    ))

    # ==============================================================================
    # SLIDE 10: NUEVO MÓDULO DE ADMINISTRACIÓN Y ASIGNACIÓN (PAMELA Y LILIANS)
    # ==============================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10, C_PERLA)
    add_header(s10, "Módulo Central de Administración: Registro y Asignación de Pacientes", "ADMINISTRACIÓN CENTRAL · PAMELA QUIJADA & LILIANS CORTÉS")

    c_adm_izq = add_card(s10, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf_adm_izq = c_adm_izq.text_frame
    tf_adm_izq.word_wrap = True
    tf_adm_izq.margin_left = tf_adm_izq.margin_right = Inches(0.35)
    tf_adm_izq.margin_top = Inches(0.25)

    p = tf_adm_izq.paragraphs[0]
    p.text = "👑 Ingreso de Paciente & Asignación en Cascada (2 Pasos)"
    p.font.name = "Georgia"
    p.font.size = Pt(14.5)
    p.font.bold = True
    p.font.color.rgb = C_GRANATE
    p.space_after = Pt(8)

    p = tf_adm_izq.add_paragraph()
    p.text = (
        "• Perfil Administrador Central (Dra. Pamela Quijada & Lilians Cortés):\n"
        "  Acceso directo en el Portal Profesional con el botón destacado '[ 👑 Asignar Paciente (Admin) ]' y pestaña de Administración.\n\n"
        "• Modo Dual: Consultante Registrado vs. Nuevo:\n"
        "  - Consultante Registrado: Selector de pacientes existentes que autocompleta instantáneamente Nombre, RUT, Correo, Teléfono y Motivo.\n"
        "  - Nuevo Consultante: Registro ágil desde cero con validación de RUT Módulo 11 (para boleta de reembolso) y dirección física.\n\n"
        "• Filtro en Cascada Inteligente (2 Pasos):\n"
        "  1. Paso 1: Selección de Disciplina de la Salud (Psicología, Nutrición, Psiquiatría, Fonoaudiología, T.O., Trabajo Social, Mediación).\n"
        "  2. Paso 2: El selector de especialistas se filtra dinámicamente mostrando solo los profesionales de esa área con su arancel.\n\n"
        "• Parámetros de Atención:\n"
        "  Definición de Fecha, Horario, Modalidad (Videoconsulta Online o Terreno) y Arancel en CLP."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_NOCHE

    c_adm_der = add_card(s10, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf_adm_der = c_adm_der.text_frame
    tf_adm_der.word_wrap = True
    tf_adm_der.margin_left = tf_adm_der.margin_right = Inches(0.35)
    tf_adm_der.margin_top = Inches(0.25)

    p = tf_adm_der.paragraphs[0]
    p.text = "📹 Sincronización Automática, WebRTC & Correos Duales"
    p.font.name = "Georgia"
    p.font.size = Pt(14.5)
    p.font.bold = True
    p.font.color.rgb = C_PETROLEO
    p.space_after = Pt(8)

    p = tf_adm_der.add_paragraph()
    p.text = (
        "• Generación de Código de Sesión Unívoco:\n"
        "  El motor asigna un código oficial trazable 'CBL-SES-2026-XXX' vinculado a la ficha clínica y a la agenda del terapeuta.\n\n"
        "• Salas de Videollamada WebRTC P2P con Roles Diferenciados:\n"
        "  - Enlace Especialista: gentia.cl/clinica-sala-espera.html?sala=CBL-SES-...&role=profesional\n"
        "  - Enlace Consultante: gentia.cl/clinica-sala-espera.html?sala=CBL-SES-...&role=paciente\n\n"
        "• Inyección Automática en Agenda del Profesional:\n"
        "  La cita aparece de inmediato en 'Próximas Sesiones' del especialista con link de videollamada y ficha clínica.\n\n"
        "• Doble Correo Corporativo Oficial Cabal Salud:\n"
        "  1. Al Especialista: Con datos del paciente, motivo, horario y link WebRTC profesional.\n"
        "  2. Al Consultante: Mensaje de bienvenida personalizado, indicaciones técnicas y link de ingreso con los colores corporativos de Cabal Salud.\n\n"
        "• Previsualización en Vivo de Correos:\n"
        "  Pestañas interactivas en pantalla para auditar el diseño HTML antes de enviar o cerrar."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(10.5)
    p.font.color.rgb = C_NOCHE

    add_slide_notes(s10, (
        "Propósito de esta lámina: Presentar el nuevo módulo exclusivo de administración central que permite a Pamela y Lilians ingresar consultantes y asignarlos a los profesionales de la clínica.\n\n"
        "Aspectos destacados para Pamela y Lilians:\n"
        "• Selector de Consultante Registrado: Si el paciente ya fue atendido o registrado antes, no es necesario volver a escribir sus datos; al seleccionarlo del desplegable, su Nombre, RUT, Correo, Teléfono y Motivo se autocompletan en 1 clic.\n"
        "• Asignación en Cascada en 2 Pasos: Primero seleccionan la Disciplina (ej. Psicología, Nutrición, Matronería, etc.) y de inmediato el listado de terapeutas se filtra mostrando solo a los especialistas de esa rama con su respectivo arancel.\n"
        "• Coordinación integral de la sesión: Pueden fijar fecha, hora, modalidad (online o terreno) y notas de recepción interna.\n"
        "• Inyección instantánea al especialista: La hora aparece de inmediato en la agenda del profesional con su enlace WebRTC y su ficha clínica vinculada.\n"
        "• Despacho de Correos Duales Corporativos: El sistema redacta y envía dos correos con el membrete y paleta de Cabal Salud: uno al terapeuta con los antecedentes clínicos y el link profesional, y otro al paciente dándole la bienvenida con indicaciones de conexión claras. Pueden previsualizar ambos correos en pantalla antes de confirmar.\n\n"
        "Feedback esperado: Les sugerimos probar la asignación de una sesión y contarnos si el orden de los campos y el tono de los correos automáticos les resultan cómodos."
    ))

    # ==============================================================================
    # SLIDE 11: PACIENTE TRANSVERSAL, TRASPASO INTERNO Y EXPEDIENTE .ZIP
    # ==============================================================================
    s11_multi = prs.slides.add_slide(blank_layout)
    set_slide_background(s11_multi, C_PERLA)
    add_header(s11_multi, "Expediente Multidisciplinario Compartido, Traspaso Interno y Descarga .ZIP", "RED MULTIDISCIPLINARIA · LEY N° 20.584 Y D.S. N° 41")

    c_multi_izq = add_card(s11_multi, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf_multi_izq = c_multi_izq.text_frame
    tf_multi_izq.word_wrap = True
    tf_multi_izq.margin_left = tf_multi_izq.margin_right = Inches(0.35)
    tf_multi_izq.margin_top = Inches(0.25)

    p = tf_multi_izq.paragraphs[0]
    p.text = "👥 Paciente Transversal Compartido (Constanza Morales)"
    p.font.name = "Georgia"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_PETROLEO
    p.space_after = Pt(8)

    p = tf_multi_izq.add_paragraph()
    p.text = (
        "• Consultante Única para Toda la Red Cabal Salud:\n"
        "  Constanza Morales Henríquez (28 años, Triatleta & Diseñadora, RUT 18.392.104-5). "
        "Al ingresar cualquier especialista a su portal, Constanza aparece en su directorio.\n\n"
        "• Fichas Ad-Hoc Precargadas para las 9 Especialidades:\n"
        "  - Psicología Deportiva: Examen Mental V70, ansiedad precompetitiva y motivo latente.\n"
        "  - Nutrición & Dietética: Evaluación ISAK II (IMC 21.7, 17.5% grasa, timing de HC).\n"
        "  - Matronería & Salud Sexual: FUR, ciclos regulares, DIU, PAP vigente, piso pélvico.\n"
        "  - Terapia Ocupacional: Perfil Sensorial SPM-2, AVD y ergonomía en teletrabajo.\n"
        "  - Fonoaudiología: Protocolo MBGR, bruxismo e hipertonía de maseteros.\n"
        "  - Trabajo Social: Diagnóstico sociofamiliar, tipología y redes de soporte.\n"
        "  - Psicología Neurodivergencia: Modulación sensorial y estilo de procesamiento.\n"
        "  - Psiquiatría: Examen somático, cronobiología y control de melatonina.\n"
        "  - Mediación Legal: Asesoría contractual de auspicio deportivo y compliance.\n\n"
        "• Trazabilidad y Confidencialidad:\n"
        "  Cada profesional edita su área manteniendo resguardo ético según la Ley N° 20.584."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(10.2)
    p.font.color.rgb = C_NOCHE

    c_multi_der = add_card(s11_multi, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf_multi_der = c_multi_der.text_frame
    tf_multi_der.word_wrap = True
    tf_multi_der.margin_left = tf_multi_der.margin_right = Inches(0.35)
    tf_multi_der.margin_top = Inches(0.25)

    p = tf_multi_der.paragraphs[0]
    p.text = "🔄 Traspaso en Cascada, Informes Oficiales & ZIP"
    p.font.name = "Georgia"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = C_ESMERALDA
    p.space_after = Pt(8)

    p = tf_multi_der.add_paragraph()
    p.text = (
        "• Sistema de Traspaso y Derivación en 2 Pasos:\n"
        "  1. Disciplina Destino: Psicología, Nutrición, Matronería, T.O., Fonoaudiología, etc.\n"
        "  2. Especialista Receptor: Selección del colega disponible en dicha área.\n"
        "  - Modalidad: Interconsulta colaborativa vs. Traspaso definitivo de custodia.\n"
        "  - Selección de Fichas Adjuntas: Anamnesis, notas DAP de evolución y Epicrisis.\n"
        "  - Sellado Criptográfico: Inserta registro inmutable con hash SHA-256 en historial.\n\n"
        "• Descarga de Informes Formateados por Especialidad:\n"
        "  Modal interactivo '[ 📄 Informe Especialidad ]' con vista previa de impresión "
        "y descarga HTML/PDF con membrete formal de Cabal Salud SpA para Isapres y peritajes.\n\n"
        "• Empaquetado Integral en Archivo .ZIP (JSZip):\n"
        "  Un clic en '[ 📦 Descargar .ZIP ]' comprime al instante sin servidores externos:\n"
        "  - 00_INDICE_EXPEDIENTE_CABAL_SALUD.txt (Índice maestro de documentos).\n"
        "  - 01_CERTIFICADO_CUSTODIA_AUDITORIA_SHA256.txt (Certificación D.S. 41 y Ley 20.584).\n"
        "  - 02_EXPEDIENTE_MAESTRO_CONSOLIDADO.html (Unificación de las 9 disciplinas).\n"
        "  - Carpeta 'Especialidades_Clinicas/' con las 9 fichas individuales formateadas."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(10.2)
    p.font.color.rgb = C_NOCHE

    add_slide_notes(s11_multi, (
        "Propósito de esta lámina: Exponer la potencia multidisciplinaria de Cabal Salud a través de un caso compartido, derivaciones internas y descarga documental en .ZIP.\n\n"
        "Detalles para Pamela y Lilians:\n"
        "• Paciente Transversal 'Constanza Morales Henríquez': En vez de tener pacientes de prueba diferentes y desarticulados, creamos una paciente única presente en los 9 profesionales. Cada terapeuta ve sus datos con la ficha clínica ad-hoc de su respectiva disciplina (ISAK II en Nutrición, Salud Sexual en Matronería, Perfil Sensorial en T.O., Bruxismo en Fonoaudiología, etc.).\n"
        "• Sistema de Traspaso Interno: Permite realizar una interconsulta o un traspaso definitivo del caso hacia otro profesional de la red, adjuntando selectivamente la anamnesis, las evoluciones DAP o la epicrisis con sellado SHA-256.\n"
        "• Informes Oficiales por Especialidad: Cada área puede emitir y descargar un informe formal listo para imprimir o enviar a Isapres y comisiones médicas.\n"
        "• Descarga Integral en Archivo .ZIP: Con el botón '[ 📦 Descargar .ZIP ]', el sistema compila en segundos un archivo comprimido que contiene el índice, las 9 fichas clínicas individuales y el Certificado de Custodia Legal bajo el D.S. N° 41.\n\n"
        "Feedback esperado: ¿Qué les parece el contenido clínico propuesto para cada especialidad y la utilidad de descargar todo el historial del paciente en un archivo .ZIP?"
    ))

    # ==============================================================================
    # SLIDE 12: SCREENSHOT PANEL DE ADMINISTRACIÓN EN GENTIA (PAMELA Y LILIANS)
    # ==============================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11, C_PERLA)
    add_header(s11, "Dashboard de Administración y Liquidaciones (Pamela y Lilians)")

    c_info11 = add_card(s11, Inches(0.8), Inches(1.8), Inches(4.8), Inches(5.1), C_BLANCO, RGBColor(220, 215, 205))
    tf11 = c_info11.text_frame
    tf11.word_wrap = True
    tf11.margin_left = tf11.margin_right = Inches(0.3)
    tf11.margin_top = Inches(0.25)

    p = tf11.paragraphs[0]
    p.text = "Ajustes Clave para Clínica Online"
    p.font.name = "Georgia"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_NOCHE
    p.space_after = Pt(8)

    p = tf11.add_paragraph()
    p.text = (
        "• Botón 'Asignar Sesión' (Actualizado):\n"
        "  Se renombró 'Reservar Box' por 'Asignar Sesión' tanto en el encabezado como en las celdas horarias.\n\n"
        "• Selector de Pacientes Existentes Integrado:\n"
        "  Permite elegir un paciente ya registrado para auto-rellenar todos sus datos, o crear uno nuevo que se almacena automáticamente en el CRM de la clínica.\n\n"
        "• Bloqueo de Opciones Físicas y Cupones:\n"
        "  Arriendo de box presencial, aseo y cupones de box quedan tachados/bloqueados por tratarse de una clínica 100% online.\n\n"
        "• Seguridad Zero-Trust:\n"
        "  Solo Pamela y Lilians pueden acceder tras iniciar sesión (se bloqueó el acceso anónimo de incógnito).\n\n"
        "• Reporte de Liquidaciones:\n"
        "  Muestra el total de sesiones atendidas por profesional para transferir sus honorarios."
    )
    p.font.name = "Calibri"
    p.font.size = Pt(11)

    img_path_5 = os.path.join(os.getcwd(), "presentation_assets", "05_gentia_dashboard.png")
    if os.path.exists(img_path_5):
        frame = add_card(s11, Inches(5.85), Inches(1.8), Inches(6.65), Inches(5.1), C_BLANCO, C_NOCHE)
        s11.shapes.add_picture(img_path_5, Inches(5.95), Inches(1.9), width=Inches(6.45), height=Inches(4.6))
        cap_box = s11.shapes.add_textbox(Inches(5.95), Inches(6.55), Inches(6.45), Inches(0.3))
        cap_tf = cap_box.text_frame
        cap_p = cap_tf.paragraphs[0]
        cap_p.text = "Acceso en vivo: https://gentia.cl/clinica-dashboard.html?id=cabalsalud"
        cap_p.font.size = Pt(10)
        cap_p.font.color.rgb = C_SLATE
        cap_p.alignment = PP_ALIGN.CENTER

    add_slide_notes(s11, (
        "Propósito de esta lámina: Mostrar el Dashboard de administración clínica en gentia.cl que Pamela y Lilians pueden utilizar para supervigilar la operación global.\n\n"
        "Ajustes aplicados recientemente:\n"
        "• Botón 'Asignar Sesión': Cambiamos la antigua etiqueta 'Reservar Box' por 'Asignar Sesión', adecuada para la atención virtual.\n"
        "• Selección de Pacientes Existentes: Se incorporó un selector de pacientes registrados que autocompleta el RUT, correo y teléfono sin tener que reescribirlos. Si ingresan un paciente nuevo, queda guardado de inmediato en su CRM para futuras atenciones.\n"
        "• Eliminación de opciones presenciales: Dado que Cabal Salud es una clínica online, dejamos desactivadas y tachadas las opciones de aseo, arriendo de box físico y cupones de arriendo.\n"
        "• Control de liquidaciones: Permite visualizar las sesiones atendidas por cada profesional para proceder con la transferencia mensual de sus honorarios.\n\n"
        "Feedback esperado: ¿Les resulta clara la visualización del calendario y el resumen de pacientes en este panel?"
    ))

    # ==============================================================================
    # SLIDE 13: MAPA CONCEPTUAL 3 — SEGURIDAD CRIPTOGRÁFICA Y WEBHOOKS HMAC
    # ==============================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12, C_PERLA)
    add_header(s12, "Mapa Conceptual 3: Seguridad Criptográfica y Sincronización HMAC")

    # LÍNEAS CONECTORAS
    add_connector_line(s12, Inches(4.4), Inches(2.7), Inches(4.7), Inches(2.7), C_PETROLEO, 2.0)
    add_connector_line(s12, Inches(8.6), Inches(2.7), Inches(8.9), Inches(2.7), C_PETROLEO, 2.0)
    add_connector_line(s12, Inches(6.65), Inches(2.9), Inches(6.65), Inches(3.6), C_GRANATE, 2.0)

    add_connector_line(s12, Inches(2.6), Inches(3.4), Inches(2.6), Inches(3.8), C_PETROLEO, 2.0)
    add_connector_line(s12, Inches(2.6), Inches(4.16), Inches(2.6), Inches(4.5), C_PETROLEO, 2.0)

    add_connector_line(s12, Inches(10.7), Inches(3.4), Inches(10.7), Inches(3.8), C_ESMERALDA, 2.0)
    add_connector_line(s12, Inches(10.7), Inches(4.16), Inches(10.7), Inches(4.5), C_ESMERALDA, 2.0)

    # Nodos
    add_concept_node(s12, Inches(0.8), Inches(2.2), Inches(3.6), Inches(1.2),
                     "Servidor Cabal Salud (Emisor)",
                     "Recibe confirmación Transbank/MP\ny construye cadena canónica de pago:\n'cabalsalud:citaId:buyOrder:monto:authCode'",
                     C_BLANCO, C_PETROLEO, C_NOCHE)

    add_link_pill(s12, Inches(4.7), Inches(2.5), Inches(3.9), Inches(0.4),
                  "firma con Secreto Compartido HMAC-SHA256", C_PERLA, C_PETROLEO)

    add_concept_node(s12, Inches(8.9), Inches(2.2), Inches(3.6), Inches(1.2),
                     "Motor Gentia (Receptor Seguro)",
                     "Valida firma en /clinica/confirmarPagoExterno.\nSi el hash coincide al 100%, activa la cita.\nSi no coincide, bloquea (401 Unauthorized)",
                     C_LIGHT_EMERALD, C_ESMERALDA, C_NOCHE)

    add_link_pill(s12, Inches(1.0), Inches(3.8), Inches(3.2), Inches(0.36),
                  "activa canal de notificación con", C_PERLA, C_PETROLEO)
    add_concept_node(s12, Inches(0.8), Inches(4.5), Inches(3.6), Inches(1.3),
                     "Telemetría de Notificaciones",
                     "Prioriza WhatsApp Business.\nSi excede cuota o falla webhook externo,\nconmuta de forma automática a Email\npara evitar citas silenciosas.",
                     C_BLANCO, C_PETROLEO, C_NOCHE)

    add_link_pill(s12, Inches(9.1), Inches(3.8), Inches(3.2), Inches(0.36),
                  "crea sesión y suma honorario a", C_PERLA, C_ESMERALDA)
    add_concept_node(s12, Inches(8.9), Inches(4.5), Inches(3.6), Inches(1.3),
                     "Sala WebRTC y Liquidación",
                     "Genera sala de teleconsulta cifrada y suma\nla sesión al informe mensual de honorarios\npara la transferencia de Pamela y Lilians.",
                     C_BLANCO, C_NOCHE, C_NOCHE)

    add_concept_node(s12, Inches(4.9), Inches(3.6), Inches(3.5), Inches(1.0),
                     "GENTIA_WEBHOOK_HMAC_SECRET",
                     "Llave simétrica privada inyectada en variables\nde entorno de Vercel (nunca expuesta al público)",
                     C_PERLA_CARD, C_GRANATE, C_GRANATE)

    add_slide_notes(s12, (
        "Propósito de esta lámina: Describir el mecanismo de seguridad y sincronización que conecta a Cabal Salud con Gentia de manera blindada.\n\n"
        "Puntos clave para Pamela y Lilians:\n"
        "• Cada vez que se procesa una asignación o un pago en línea, Cabal Salud genera una firma criptográfica con un secreto privado (HMAC-SHA256).\n"
        "• El servidor de Gentia verifica esta firma matemática en milisegundos. Si alguien intentara manipular la información desde el exterior, la solicitud es rechazada de inmediato.\n"
        "• Telemetría con respaldo: El sistema cuenta con redundancia; si un canal de mensajería externa presenta lentitud, conmuta automáticamente a correo corporativo para garantizar que el profesional y el paciente reciban siempre su notificación.\n"
        "• Videoconsulta P2P (WebRTC): Las sesiones online se establecen directamente entre el navegador del profesional y el del paciente con cifrado de extremo a extremo.\n\n"
        "Feedback esperado: Esta lámina les sirve como respaldo técnico de alto nivel ante cualquier solicitud de seguridad de la información."
    ))

    # ==============================================================================
    # SLIDE 14: OPORTUNIDADES ESTRATÉGICAS Y RETORNO DE INVERSIÓN (ROI)
    # ==============================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13, C_PERLA)
    add_header(s13, "Oportunidades Estratégicas y Hoja de Ruta de Crecimiento")

    opps = [
        ("📱 Recordatorios WhatsApp", "Costo Meta: ~$40 CLP/msg (~$4.500/mes). Reduce el no-show de 18% a menos de 4%, rescatando +$350.000 CLP mensuales en consultas que antes se perdían por inasistencia.", C_ESMERALDA),
        ("📦 Packs de 4 Sesiones", "Descuento 10% por pago anticipado. Asegura flujo de caja inmediato ($126.000 en vez de arriesgar deserción tras sesión 1) y mejora adherencia clínica.", C_PETROLEO),
        ("🧭 Triage Clínico Guiado", "Widget de 3 preguntas en el Hero: '¿No sabes qué especialista necesitas?'. Orienta a pacientes indecisos y eleva la conversión del sitio entre 20% y 30%.", C_GRANATE),
        ("🤝 Convenios B2B Empresas", "Planes de salud mental y bienestar para colaboradores de empresas chilenas con factura exenta de IVA deducible de gastos corporativos.", C_NOCHE)
    ]

    for idx, (op_title, op_desc, op_color) in enumerate(opps):
        row = idx // 2
        col = idx % 2
        left_x = Inches(0.8 + col * 6.0)
        top_y = Inches(1.9 + row * 2.6)
        
        card = add_card(s13, left_x, top_y, Inches(5.7), Inches(2.35), C_BLANCO, op_color)
        ctf = card.text_frame
        ctf.word_wrap = True
        ctf.margin_left = ctf.margin_right = Inches(0.3)
        ctf.margin_top = Inches(0.2)
        
        p = ctf.paragraphs[0]
        p.text = op_title
        p.font.name = "Georgia"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = op_color
        p.space_after = Pt(6)
        
        p = ctf.add_paragraph()
        p.text = op_desc
        p.font.name = "Calibri"
        p.font.size = Pt(11.5)
        p.font.color.rgb = C_NOCHE

    add_slide_notes(s13, (
        "Propósito de esta lámina: Presentar 4 oportunidades estratégicas de negocio que pueden evaluar e implementar cuando lo consideren oportuno.\n\n"
        "Opciones presentadas:\n"
        "1. Recordatorios por WhatsApp Business: Automatizar avisos 24 horas y 2 horas antes de la sesión. Es la medida más efectiva para reducir el ausentismo (no-show) de un 18% a menos de un 4%, protegiendo los ingresos de los profesionales.\n"
        "2. Venta de Packs de Continuidad (4 sesiones): Ofrecer un paquete mensual con un pequeño descuento incentiva a que el paciente complete su proceso terapéutico y brinda liquidez inicial a la clínica.\n"
        "3. Triage Clínico Guiado: Un breve cuestionario orientador en la página principal para quienes no saben si consultar con psicología, nutrición o psiquiatría, facilitando la decisión de reserva.\n"
        "4. Convenios B2B para Empresas: Programas de bienestar emocional para organizaciones con emisión de facturas exentas de IVA, un canal de alto volumen y recurrencia.\n\n"
        "Feedback esperado: ¿Cuál de estas 4 iniciativas les parece más atractiva para priorizar en la siguiente etapa de Cabal Salud?"
    ))

    # ==============================================================================
    # SLIDE 15: CONCLUSIONES Y ARGUMENTARIO ANTE AUDITORÍAS
    # ==============================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14, C_NOCHE)
    add_header(s14, "Argumentario de Defensa ante Auditorías y Enlaces Productivos", "BLINDAJE COMPLETO · CABAL SALUD")

    c_fin = add_card(s14, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.1), RGBColor(22, 48, 75), C_PETROLEO)
    tf_fin = c_fin.text_frame
    tf_fin.word_wrap = True
    tf_fin.margin_left = tf_fin.margin_right = Inches(0.4)
    tf_fin.margin_top = Inches(0.3)

    p = tf_fin.paragraphs[0]
    p.text = "Argumentario Oficial para Pamela Quijada y Lilians Cortés ante Fiscalizaciones"
    p.font.name = "Georgia"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = C_PERLA
    p.space_after = Pt(10)

    p = tf_fin.add_paragraph()
    p.text = (
        "1. Frente a una Auditoría Tributaria del SII:\n"
        "   'Clínica Cabal Salud SpA recauda el 100% de los honorarios directamente en sus cuentas bancarias "
        "mediante Transbank Webpay Plus y Mercado Pago propios. Las prestaciones médicas están exentas de IVA (Art. 12 E N° 8) "
        "y no existe intermediación ni recaudación por terceros.'\n\n"
        "2. Frente a una Auditoría Sanitaria (Superintendencia de Salud / Ley 20.584 y D.S. N° 41):\n"
        "   'El sitio web comercial cabalsalud.cl no almacena fichas clínicas ni datos de salud sensibles. "
        "Cabal Salud mantiene un contrato de custodia e infraestructura con Gentia.cl, donde las historias clínicas y evoluciones "
        "son almacenadas bajo cifrado AES-256 en reposo, con firma inmutable y retención legal de 15 años.'\n\n"
        "3. Enlaces Oficiales de Operación en Producción:\n"
        "   • Sitio Pacientes: https://cabalsalud.vercel.app  (o https://cabalsalud.cl al delegar DNS)\n"
        "   • Portal Especialistas (Fichas, Arancel y Tiempo 30-90m): https://cabalsalud.vercel.app/portal-profesional.html\n"
        "   • Dashboard Pamela y Lilians: https://gentia.cl/clinica-dashboard.html?id=cabalsalud"
    )
    p.font.name = "Calibri"
    p.font.size = Pt(11.5)
    p.font.color.rgb = C_BLANCO

    add_slide_notes(s14, (
        "Propósito de esta lámina: Brindarles las respuestas oficiales listas para responder con solidez ante cualquier requerimiento del SII o de la Superintendencia de Salud, y facilitarles los accesos directos.\n\n"
        "Mensaje de cierre para Pamela y Lilians:\n"
        "• Tienen en sus manos un ecosistema que combina calidez de atención con rigurosidad legal, informática y contable de primer nivel.\n"
        "• Los enlaces para ingresar y explorar en vivo son:\n"
        "  - Sitio público de pacientes: https://cabalsalud.vercel.app\n"
        "  - Portal de profesionales y administración: https://cabalsalud.vercel.app/portal-profesional.html\n"
        "  - Dashboard de gestión en Gentia: https://gentia.cl/clinica-dashboard.html?id=cabalsalud\n\n"
        "Queremos reiterarles que no hay ningún apuro. Revisen todo con calma, naveguen por la plataforma y prueben las funciones. Estaremos atentos a todos sus comentarios, sugerencias y reparos (tanto de forma estética como de fondo operativo) para dejárselos exactamente como lo soñaron."
    ))

    # Guardar y actualizar archivos
    output_pptx_workspace = os.path.join(r"C:\Users\mtras\Desktop\Proyectos varios\cabalsalud", "Presentacion_Cabalsalud_Arquitectura_y_Funcionamiento.pptx")
    prs.save(output_pptx_workspace)
    print(f"Presentación actualizada exitosamente en: {output_pptx_workspace}")

    brain_dir = r"C:\Users\mtras\.gemini\antigravity\brain\03572d44-b3c4-46c2-b2a0-eb2585e5b2eb"
    if os.path.exists(brain_dir):
        output_pptx_brain = os.path.join(brain_dir, "Presentacion_Cabalsalud_Arquitectura_y_Funcionamiento.pptx")
        shutil.copyfile(output_pptx_workspace, output_pptx_brain)
        print(f"Copia actualizada en artefactos: {output_pptx_brain}")

if __name__ == "__main__":
    generar_presentacion_completa()
